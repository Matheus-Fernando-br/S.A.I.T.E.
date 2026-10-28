"""Extração de texto de arquivos, com validação do tipo real do conteúdo.

Cada extrator devolve uma lista de ``Section`` (texto + localização legível,
como "p. 3" ou "Seção: Introdução"), usada depois nas citações.
"""

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import PurePath


class ExtractionError(ValueError):
    pass


@dataclass
class Section:
    text: str
    location: str = ""


SUPPORTED = {
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".csv": "text/csv",
    ".html": "text/html",
    ".htm": "text/html",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

# Limites contra arquivos compactados maliciosos (DOCX/XLSX são ZIP).
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024
MAX_ZIP_RATIO = 100
MAX_PDF_PAGES = 2000
MAX_SHEET_ROWS = 50_000

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def normalize_text(text: str) -> str:
    """NFC, remove caracteres de controle e espaços redundantes, preservando parágrafos."""
    text = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL.sub("", text).replace(" ", " ")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def detect_type(filename: str, data: bytes) -> str:
    """Retorna a extensão validada ou levanta ``ExtractionError``."""
    suffix = PurePath(filename).suffix.lower()
    if suffix not in SUPPORTED:
        allowed = ", ".join(sorted({s.lstrip(".").upper() for s in SUPPORTED}))
        raise ExtractionError(f"Tipo de arquivo não suportado ({suffix or 'sem extensão'}). Aceitos: {allowed}.")
    if not data:
        raise ExtractionError("Arquivo vazio.")
    if suffix == ".pdf" and not data.startswith(b"%PDF"):
        raise ExtractionError("O conteúdo não é um PDF válido.")
    if suffix in {".docx", ".xlsx"}:
        if not data.startswith(b"PK\x03\x04"):
            raise ExtractionError(f"O conteúdo não é um {suffix.upper()[1:]} válido.")
        _check_zip(data)
    if suffix in {".txt", ".md", ".markdown", ".csv", ".html", ".htm"} and b"\x00" in data[:8192]:
        raise ExtractionError("O arquivo parece ser binário, não texto.")
    return suffix


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ExtractionError("Arquivo compactado corrompido.") from exc
    if total > MAX_UNZIPPED_BYTES or total > len(data) * MAX_ZIP_RATIO:
        raise ExtractionError("Arquivo recusado: conteúdo descompactado grande demais.")


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Não foi possível decodificar o texto.")


def extract(filename: str, data: bytes) -> list[Section]:
    suffix = detect_type(filename, data)
    extractor = {
        ".pdf": _pdf,
        ".docx": _docx,
        ".xlsx": _xlsx,
        ".csv": _csv,
        ".html": _html_bytes,
        ".htm": _html_bytes,
        ".md": _markdown,
        ".markdown": _markdown,
        ".txt": _plain,
    }[suffix]
    try:
        sections = extractor(data)
    except ExtractionError:
        raise
    except Exception as exc:  # bibliotecas de terceiros levantam tipos variados
        raise ExtractionError(f"Falha ao ler o arquivo: {type(exc).__name__}: {exc}") from exc
    sections = [Section(normalize_text(s.text), s.location) for s in sections]
    sections = [s for s in sections if s.text]
    if not sections:
        if suffix == ".pdf":
            raise ExtractionError(
                "Nenhum texto encontrado no PDF. Ele pode ser digitalizado (imagem); "
                "OCR ainda não está disponível — exporte o documento com texto ou use outro formato."
            )
        raise ExtractionError("Nenhum texto encontrado no arquivo.")
    return sections


def _plain(data: bytes) -> list[Section]:
    return [Section(decode_text(data))]


def _markdown(data: bytes) -> list[Section]:
    return split_by_headings(decode_text(data).splitlines(), r"^(#{1,3})\s+(.+?)\s*#*$")


def split_by_headings(lines: list[str], pattern: str) -> list[Section]:
    heading = re.compile(pattern)
    sections: list[Section] = []
    current: list[str] = []
    title = ""
    for line in lines:
        match = heading.match(line)
        if match:
            if "\n".join(current).strip():
                sections.append(Section("\n".join(current), f"Seção: {title}" if title else ""))
            title = match.group(2).strip()
            current = [line]
        else:
            current.append(line)
    if "\n".join(current).strip():
        sections.append(Section("\n".join(current), f"Seção: {title}" if title else ""))
    return sections


def _pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("PDF protegido por senha não é suportado.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ExtractionError(f"PDF com mais de {MAX_PDF_PAGES} páginas.")
        return [
            Section(page.extract_text() or "", f"p. {number}")
            for number, page in enumerate(reader.pages, start=1)
        ]
    except PdfReadError as exc:
        raise ExtractionError(f"PDF inválido: {exc}") from exc


def _docx(data: bytes) -> list[Section]:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        text = paragraph.text
        if style.lower().startswith(("heading", "título", "titulo")) and text.strip():
            lines.append(f"# {text}")
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return split_by_headings(lines, r"^(#)\s+(.+)$")


def _xlsx(data: bytes) -> list[Section]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_SHEET_ROWS:
                    break
                cells = ["" if value is None else str(value) for value in row]
                if any(cell.strip() for cell in cells):
                    rows.append(" | ".join(cells).rstrip(" |"))
            if rows:
                sections.append(Section("\n".join(rows), f"Planilha: {sheet.title}"))
    finally:
        workbook.close()
    return sections


def _csv(data: bytes) -> list[Section]:
    text = decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))[:MAX_SHEET_ROWS]
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    sections = []
    # Grupos de linhas repetindo o cabeçalho, para que cada trecho seja autoexplicativo.
    for start in range(0, max(len(body), 1), 50):
        group = body[start : start + 50]
        lines = [" | ".join(header)] + [" | ".join(row) for row in group]
        sections.append(Section("\n".join(lines), f"linhas {start + 2}–{start + 1 + len(group)}"))
    return sections


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "template", "svg", "nav", "footer", "form", "iframe"}
    BLOCK = {"p", "div", "br", "li", "tr", "section", "article", "h1", "h2", "h3", "h4", "pre", "table"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.lines: list[str] = []
        self.buffer: list[str] = []
        self.skip_depth = 0
        self.heading: str | None = None
        self.title = ""
        self.in_title = False

    def flush(self):
        text = "".join(self.buffer).strip()
        if text:
            self.lines.append(f"# {text}" if self.heading else text)
        self.buffer = []

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip_depth += 1
        elif tag == "title":
            self.in_title = True
        elif tag in self.BLOCK:
            self.flush()
            if tag in {"h1", "h2", "h3"}:
                self.heading = tag

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip_depth = max(0, self.skip_depth - 1)
        elif tag == "title":
            self.in_title = False
        elif tag in self.BLOCK:
            self.flush()
            if tag in {"h1", "h2", "h3"}:
                self.heading = None

    def handle_data(self, data):
        if self.in_title:
            self.title += data
        elif not self.skip_depth:
            self.buffer.append(data)


def html_to_sections(html: str) -> tuple[str, list[Section]]:
    """Converte HTML em seções por título. Retorna (título da página, seções)."""
    parser = _HTMLText()
    parser.feed(html)
    parser.close()
    parser.flush()
    return parser.title.strip(), split_by_headings(parser.lines, r"^(#)\s+(.+)$")


def _html_bytes(data: bytes) -> list[Section]:
    return html_to_sections(decode_text(data))[1]
