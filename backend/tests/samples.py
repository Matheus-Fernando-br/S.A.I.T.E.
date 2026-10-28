"""Gera arquivos de exemplo reais (PDF, DOCX, XLSX) para os testes de ingestão."""

import io


def make_pdf(pages: list[str]) -> bytes:
    """PDF mínimo válido, uma linha de texto por página (fonte Helvetica)."""
    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(len(pages))]
    font_id = 3 + 2 * len(pages)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for index, text in enumerate(pages):
        content_id = page_ids[index] + 1
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R"
            f" /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx(heading: str, paragraphs: list[str]) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Estoque"
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
